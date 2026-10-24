# -*- coding: utf-8 -*-
"""整体核对：页数/页码/禁词/密度/表格/图片"""
from pptx import Presentation
from pptx.util import Emu
import re, sys

PATH = r"C:\Users\Su\Desktop\完美假日\车记得复赛汇报\车记得复赛汇报.pptx"
prs = Presentation(PATH)
slides = list(prs.slides)
print(f"总页数: {len(slides)} (期望 12)")
print(f"画布: {prs.slide_width}x{prs.slide_height} EMU = {prs.slide_width/914400*72:.0f}x{prs.slide_height/914400*72:.0f} pt")

def shape_texts(shp):
    out = []
    if shp.has_text_frame:
        for p in shp.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs)
            if t.strip():
                out.append(t)
    if shp.shape_type == 6:  # group
        for s in shp.shapes:
            out += shape_texts(s)
    if getattr(shp, "has_table", False) and shp.has_table:
        for row in shp.table.rows:
            for c in row.cells:
                if c.text.strip():
                    out.append(c.text)
    return out

def count_pics(shp):
    n = 0
    if shp.shape_type == 13:
        n += 1
    if shp.shape_type == 6:
        for s in shp.shapes:
            n += count_pics(s)
    return n

all_text = []
issues = []
for i, sl in enumerate(slides, 1):
    texts, pics, tables = [], 0, 0
    for shp in sl.shapes:
        texts += shape_texts(shp)
        pics += count_pics(shp)
        if getattr(shp, "has_table", False) and shp.has_table:
            tables += 1
            rows = len(shp.table.rows); cols = len(shp.table.columns)
            print(f"  P{i:02d} 表格: {rows}行 x {cols}列")
    body = " ".join(texts)
    all_text.append(body)
    # 页码核对
    m = re.search(r"(\d{2})\s*/\s*12", body)
    page_ok = m and int(m.group(1)) == i
    # 中文字数
    cn = len(re.findall(r"[一-鿿]", body))
    # 母版：标题/页脚
    has_footer = "车记得 · 完美金鹰假日商圈 AI 大赛" in body
    flag = "" if (page_ok and has_footer) else "  <-- 缺页码或页脚"
    if i == 1:
        flag = "  (封面按契约无页码/页脚)"
        page_ok, has_footer = True, True
    print(f"P{i:02d}: 文本块{len(texts):3d} | 中文{cn:3d}字 | 图{pics} | 表{tables} | 页码{'OK' if page_ok else 'MISS'} | 页脚{'OK' if has_footer else 'MISS'}{flag}")
    if not page_ok:
        issues.append(f"P{i} 页码缺失或不符")
    if i not in (1,) and not has_footer:
        issues.append(f"P{i} 页脚缺失")

full = "\n".join(all_text)
# 数据纪律禁词
for bad, why in [("4500", "未核实车位数"), ("8 分钟", "未实测寻车时长"), ("3000 万", "未核实客流"), ("8分钟", "未实测寻车时长"), ("3000万", "未核实客流")]:
    if bad in full:
        issues.append(f"禁词命中: {bad} ({why})")
# 必须标注
musts = [
    ("非实测", "demo 模拟数据标注", [6, 12]),
    ("实地调研", "调研来源标注", [3]),
    ("待补充", "团队占位", [9]),
    ("关键词模拟", "AI 模拟口径", [6, 12]),
]
for kw, why, pages in musts:
    for p in pages:
        if kw not in all_text[p-1]:
            issues.append(f"P{p} 缺少必须标注: {kw} ({why})")

# 色板纪律（slide 源文件层面另查）
print()
if issues:
    print("!! 发现问题:")
    for x in issues:
        print("  -", x)
    sys.exit(1)
print("结构核对全部通过")
